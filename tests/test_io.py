"""Belge okuyucular ve rapor dışa aktarıcılar."""

from __future__ import annotations

import io
import json
import re
from pathlib import Path

import docx
import pytest
from pypdf import PdfWriter

from kolaymetin import analyze
from kolaymetin.io.exporters import dither_level, render_marked
from kolaymetin.io.readers import ReaderError, read_bytes, read_file

TEXT = (
    "SU KESİNTİSİ HAKKINDA DUYURU\n\nMüracaatlar 15.09.2026 tarihinde ivedilikle alınacaktır; "
    "aksi takdirde başvurmamak mümkün değildir.\n\nSu gelecek."
)


def _text_pdf(text: str) -> bytes:
    """Metin katmanı olan küçük bir PDF (yalnızca test için, elle yazılmış)."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, o in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def test_read_txt_encodings(tmp_path: Path) -> None:
    p = tmp_path / "a.txt"
    p.write_bytes("Şu ığdır".encode("cp1254"))
    assert read_file(p) == "Şu ığdır"
    p.write_bytes("﻿Şu".encode())
    assert read_file(p) == "Şu"
    assert read_bytes("# Başlık".encode(), "notlar.md") == "# Başlık"


def test_read_docx_with_list_and_table() -> None:
    d = docx.Document()
    d.add_paragraph("Başlık")
    d.add_paragraph("Birinci madde", style="List Bullet")
    d.add_paragraph("İkinci madde", style="List Bullet")
    d.add_paragraph("")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Ad"
    table.rows[0].cells[1].text = "Ali"
    buf = io.BytesIO()
    d.save(buf)
    text = read_bytes(buf.getvalue(), "belge.docx")
    assert "Başlık\n\n- Birinci madde\n- İkinci madde" in text
    assert "Ad – Ali" in text


def test_read_pdf_text_and_scanned() -> None:
    assert "Su kesilecek" in read_bytes(_text_pdf("Su kesilecek"), "d.pdf")
    blank = PdfWriter()
    blank.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    blank.write(buf)
    with pytest.raises(ReaderError, match="metin bulunamadı"):
        read_bytes(buf.getvalue(), "tarama.pdf")


@pytest.mark.parametrize(
    ("data", "name", "msg"),
    [(b"x", "a.exe", "desteklenmiyor"), (b"bozuk", "a.docx", "açılamadı"), (b"bozuk", "a.pdf", "açılamadı")],
)
def test_reader_errors(data: bytes, name: str, msg: str) -> None:
    with pytest.raises(ReaderError, match=msg):
        read_bytes(data, name)


def test_read_file_errors(tmp_path: Path) -> None:
    with pytest.raises(ReaderError, match="bulunamadı"):
        read_file(tmp_path / "yok.txt")
    with pytest.raises(ReaderError):
        read_file(tmp_path)


def test_json_roundtrip() -> None:
    r = analyze(TEXT)
    data = json.loads(r.to_json())
    assert data["version"] == "1.0.0" and data["findings"]
    assert data["sentences"][0]["is_heading"] is True
    assert r.to_dict()["profile"] == "kolay-dil"


def test_markdown_report() -> None:
    md = analyze(TEXT).to_markdown()
    for part in ("# kolaymetin denetim raporu", "Kolay Dil Uyum Skoru", "Bu skor nasıl hesaplandı?",
                 "■ Hata", "→ Öneri", "Rehber: /rehber/"):
        assert part in md
    assert "Bu profilde bulgu yok." in analyze("Su gelecek.").to_markdown()


def test_html_report_is_standalone_and_follows_style() -> None:
    html = analyze(TEXT).to_html()
    assert html.startswith("<!DOCTYPE html>") and 'lang="tr"' in html
    assert "imi--hata" in html and "not-kagidi" in html and "d-00" in html
    # Bağımsız dosya: hiçbir dış kaynak yok, yazı tipleri ve desenler gömülü.
    assert not re.search(r"""(src|href)=["']?(https?:)?//""", html)
    assert not re.search(r"url\((['\"])?(https?:)?//", html)
    assert "data:font/woff2;base64," in html and "data:image/png;base64," in html
    assert "gradient" not in html.lower()
    assert "@page" in html and "18mm" in html
    # Künyede dörtlü: renk style.css'ten gelir, isaret.svg'nin kendi stili taşınmaz. Sekme simgesi gömülü.
    kunye = html[html.index('<header class="kunye">'):html.index("</header>")]
    assert '<svg class="dortlu"' in kunye and kunye.count('class="dortlu__im ') == 3
    assert "<style>" not in kunye and "#B42D26" not in kunye
    assert '<link rel="icon" href="data:image/svg+xml;base64,' in html


def test_render_marked_overlaps() -> None:
    r = analyze("Müracaatlarınızı ivedilikle yapın ve belgeleri getirin, sonra bekleyin.")
    marked = render_marked(r.text, r.findings, 0, {i: i + 1 for i in range(len(r.findings))})
    assert "imi--uyari" in marked and "imi-no" in marked
    assert render_marked("abc", [], 0) == "abc"


def test_dither_level() -> None:
    assert dither_level(0) == 0 and dither_level(100) == 16 and dither_level(50) == 8
    assert dither_level(-5) == 0 and dither_level(150) == 16
